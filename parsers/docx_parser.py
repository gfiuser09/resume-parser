"""DOCX text extraction using python-docx.

Tables are included because resumes frequently lay out skills, dates and
institutions inside invisible tables.
"""

import logging

from utils.errors import CorruptedFile, TextExtractionError

logger = logging.getLogger(__name__)

try:
    import docx  # python-docx
except ImportError:  # pragma: no cover
    docx = None


def _table_lines(table):
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        # de-duplicate merged cells that repeat the same text
        unique = []
        for cell in cells:
            if cell and (not unique or unique[-1] != cell):
                unique.append(cell)
        if unique:
            yield " | ".join(unique)


def extract_text(path):
    """Return the plain text of the DOCX file at *path*."""
    if docx is None:  # pragma: no cover
        raise TextExtractionError(
            "python-docx is not installed. Run: pip install python-docx", status_code=503
        )

    try:
        document = docx.Document(path)
    except Exception as exc:
        logger.warning("failed to open DOCX: %s", exc)
        raise CorruptedFile(
            "Could not open the DOCX file. It may be corrupted, or it may be a "
            "legacy .doc file - re-save the resume as .docx"
        ) from exc

    parts = [paragraph.text.strip() for paragraph in document.paragraphs]

    for table in document.tables:
        parts.extend(_table_lines(table))

    text = "\n".join(part for part in parts if part is not None).strip()
    if not text:
        raise TextExtractionError("The DOCX file contains no readable text")
    return text
